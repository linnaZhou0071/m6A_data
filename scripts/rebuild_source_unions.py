#!/usr/bin/env python3
"""Rebuild HEK/HeLa union and >=80% tables from the cited source files.

Default mode is read-only verification against the versioned tables.
--output-dir writes normalized CSV copies to a new directory; it refuses
to overwrite existing files. It does not copy third-party source files.
"""
from __future__ import annotations
import argparse
import csv
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd

from verify_sac_glori_sources import dec, sac_sheet

ROOT = Path(__file__).resolve().parents[1]
TOLERANCE = Decimal("0.000000001")


def formatted(value: Decimal | None) -> str:
    return "" if value is None else format(value, "f")


def gather_sources() -> dict[str, dict[str, dict[str, Decimal]]]:
    sac = openpyxl.load_workbook(ROOT / "SAC-seq_data.xlsx", read_only=True, data_only=True)
    poly, _ = sac_sheet(sac, "HEK293 polyA")
    ribo, _ = sac_sheet(sac, "HEK293 ribo-")
    hela_sheet, _ = sac_sheet(sac, "HeLa polyA")
    if any(len(values) != 1 for sheet in (poly, ribo, hela_sheet)
           for values in sheet.values()):
        raise ValueError("Unexpected duplicate position within a SAC sheet")
    hek_sac = defaultdict(list)
    for source in (poly, ribo):
        for site, values in source.items():
            hek_sac[site].extend(values)
    hek_sac_mean = {site: sum(values) / len(values) for site, values in hek_sac.items()}
    hela_sac = {site: values[0] for site, values in hela_sheet.items()}
    glori = openpyxl.load_workbook(
        ROOT / "GLORI_HEK293T_mRNA.xlsx", read_only=True, data_only=True
    )
    glori_map = {}
    rows = glori.active.iter_rows(values_only=True)
    next(rows)
    for chrom, position, strand, _gene, _cov1, _cov2, rep1, rep2, *_ in rows:
        site = f"{chrom}_{int(position)}_{strand}"
        if site in glori_map:
            raise ValueError(f"Duplicate GLORI position: {site}")
        a, b = dec(rep1), dec(rep2)
        if not (Decimal(0) <= a <= 1 and Decimal(0) <= b <= 1):
            raise ValueError(f"GLORI fraction out of range: {site}")
        glori_map[site] = (a + b) * 50
    etam = pd.read_csv(
        ROOT / "eTAM-seq data/GSE211303_hela.polya.wt.ftom.ftop.rep1.deep.hits.txt",
        sep="\t", usecols=["pos", "methylation"], dtype=str, keep_default_na=False,
    )
    if etam.pos.duplicated().any():
        raise ValueError("Duplicate eTAM position")
    etam_map = dict(zip(etam.pos, etam.methylation.map(dec)))
    return {
        "HEK": {"m6a_file1_s23": hek_sac_mean, "m6a_file2": glori_map},
        "HeLa": {"m6a_file1_sheet1": hela_sac, "m6a_file3": etam_map},
    }


def build_rows(sources: dict[str, dict[str, Decimal]]) -> list[dict[str, str]]:
    columns = list(sources)
    positions = sorted(set().union(*(set(source) for source in sources.values())))
    result = []
    for site in positions:
        chrom, position, strand = site.rsplit("_", 2)
        values = [sources[column][site] for column in columns if site in sources[column]]
        if not values:
            raise AssertionError("Site without any source")
        mean = sum(values) / len(values)
        row = {"pos": site}
        for column in columns:
            value = sources[column].get(site)
            row[column] = formatted(value)
        for column in columns:
            row[column.replace("m6a_", "in_")] = (
                "TRUE" if site in sources[column] else "FALSE"
            )
        row.update({
            "m6a_mean": formatted(mean), "chrom": chrom,
            "posnum": position, "strand": strand,
        })
        result.append(row)
    return result


def verify(name: str, rows: list[dict[str, str]], high: list[dict[str, str]]) -> None:
    for filename, expected_rows in (
        (f"{name}_union.csv", rows), (f"{name}_over_80%.csv", high)
    ):
        actual = pd.read_csv(
            ROOT / "data/processed" / filename, dtype=str, keep_default_na=False
        )
        if actual.pos.eq("").any() or actual.pos.duplicated().any():
            raise AssertionError(f"{filename}: duplicate/missing site ID")
        expected = {row["pos"]: row for row in expected_rows}
        observed = actual.set_index("pos").to_dict(orient="index")
        if set(expected) != set(observed):
            raise AssertionError(
                f"{filename}: position sets differ; missing={len(set(expected)-set(observed))}, "
                f"extra={len(set(observed)-set(expected))}"
            )
        for site, exp in expected.items():
            obs = observed[site]
            for column in actual.columns:
                if column == "pos":
                    continue
                left, right = exp[column], obs[column]
                if column.startswith("m6a_"):
                    if bool(left) != bool(right) or (
                        left and abs(dec(left) - dec(right)) > TOLERANCE
                    ):
                        raise AssertionError(f"{filename}: {site}/{column} differs")
                elif column.startswith("in_"):
                    if left.lower() != right.lower():
                        raise AssertionError(f"{filename}: {site}/{column} differs")
                elif left != right:
                    raise AssertionError(f"{filename}: {site}/{column} differs")
        print(f"PASS {filename}: {len(expected)} positions, all values/flags match")


def save(name: str, rows: list[dict[str, str]], high: list[dict[str, str]], out: Path) -> None:
    for filename, values in ((f"{name}_union.csv", rows), (f"{name}_over_80%.csv", high)):
        target = out / filename
        if target.exists():
            raise FileExistsError(target)
        with target.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=list(values[0]))
            writer.writeheader()
            writer.writerows(values)
        print(f"Wrote normalized copy {target}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path,
                        help="Optional NEW directory for normalized CSV copies")
    args = parser.parse_args()
    groups = gather_sources()
    for name, sources in groups.items():
        rows = build_rows(sources)
        high = [row for row in rows if dec(row["m6a_mean"]) >= 80]
        verify(name, rows, high)
        if args.output_dir:
            args.output_dir.mkdir(parents=True, exist_ok=True)
            save(name, rows, high, args.output_dir)


if __name__ == "__main__":
    main()
