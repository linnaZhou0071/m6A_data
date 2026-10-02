#!/usr/bin/env python3
"""Verify every SAC-seq/GLORI measurement used in HEK and HeLa union tables.

This is a read-only source-to-column audit, not a license determination.
Requires the locally downloaded supplementary workbooks.
"""
from collections import defaultdict
from decimal import Decimal
from pathlib import Path

import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]


def dec(value) -> Decimal:
    result = Decimal(str(value))
    if not result.is_finite():
        raise ValueError(f"Non-finite methylation value: {value}")
    return result


def sac_sheet(workbook, name: str):
    mapping = defaultdict(list)
    rows = 0
    for chrom, pos, strand, fraction, *_ in list(workbook[name].values)[1:]:
        site = f"{chrom}_{int(pos)}_{strand}"
        value = dec(fraction)
        if not Decimal(0) <= value <= Decimal(100):
            raise ValueError(f"SAC percentage out of range at {site}")
        mapping[site].append(value)
        rows += 1
    return mapping, rows


def source_column(union: pd.DataFrame, column: str) -> dict[str, Decimal]:
    subset = union.loc[union[column].ne(""), ["pos", column]]
    if subset.pos.duplicated().any():
        raise AssertionError(f"Duplicate union site in {column}")
    return dict(zip(subset.pos, subset[column].map(dec)))


def compare(label: str, expected: dict[str, Decimal], observed: dict[str, Decimal]) -> None:
    missing = set(expected) - set(observed)
    extra = set(observed) - set(expected)
    # Historical CSV serialization contains binary-float tails (e.g. 55.673049999999996).
    tolerance = Decimal("0.000000001")
    different = {site for site in expected.keys() & observed.keys()
                 if abs(expected[site] - observed[site]) > tolerance}
    print(f"{label}: source={len(expected)}, union={len(observed)}, "
          f"missing={len(missing)}, extra={len(extra)}, unequal={len(different)}")
    if missing or extra or different:
        print("examples:", sorted(missing)[:3], sorted(extra)[:3], sorted(different)[:3])
        raise AssertionError(f"{label} does not match source")



def verify_union_means(group: str, frame: pd.DataFrame, columns: list[str]) -> None:
    total_support = {1: 0, 2: 0}
    high_support = {1: 0, 2: 0}
    seen = set()
    for row in frame.itertuples(index=False):
        values = row._asdict()
        site = values["pos"]
        if site in seen:
            raise AssertionError(f"{group}: duplicate site {site}")
        seen.add(site)
        parts = site.rsplit("_", 2)
        if len(parts) != 3 or parts != [values["chrom"], values["posnum"], values["strand"]]:
            raise AssertionError(f"{group}: coordinate fields differ from {site}")
        method_values = []
        for column in columns:
            value = values[column]
            flag = values[column.replace("m6a_", "in_")].lower()
            if flag not in ("true", "false") or (flag == "true") != bool(value):
                raise AssertionError(f"{group}: presence flag mismatch at {site}/{column}")
            if value:
                parsed = dec(value)
                if not Decimal(0) <= parsed <= Decimal(100):
                    raise AssertionError(f"{group}: percentage out of range at {site}")
                method_values.append(parsed)
        if not method_values:
            raise AssertionError(f"{group}: no method value at {site}")
        mean = sum(method_values) / len(method_values)
        if abs(mean - dec(values["m6a_mean"])) > Decimal("0.000000001"):
            raise AssertionError(f"{group}: mean mismatch at {site}")
        n = len(method_values)
        total_support[n] += 1
        if dec(values["m6a_mean"]) >= 80:
            high_support[n] += 1
    print(f"{group} group means/flags/coordinates PASS: rows={len(frame)}, "
          f"method support={total_support}, >=80 support={high_support}")


def main() -> None:
    sac = openpyxl.load_workbook(ROOT / "SAC-seq_data.xlsx", read_only=True, data_only=True)
    hela_sac, hela_rows = sac_sheet(sac, "HeLa polyA")
    poly, poly_rows = sac_sheet(sac, "HEK293 polyA")
    ribo, ribo_rows = sac_sheet(sac, "HEK293 ribo-")
    if any(len(values) != 1 for values in hela_sac.values()) or any(
        len(values) != 1 for values in poly.values()
    ) or any(len(values) != 1 for values in ribo.values()):
        raise AssertionError("Unexpected duplicate within a SAC sheet")
    hek_sac = defaultdict(list)
    for source in (poly, ribo):
        for site, values in source.items():
            hek_sac[site].extend(values)
    hek_expected = {site: sum(values) / len(values) for site, values in hek_sac.items()}
    hela_expected = {site: values[0] for site, values in hela_sac.items()}
    hek_union = pd.read_csv(ROOT / "data/processed/HEK_union.csv", dtype=str, keep_default_na=False)
    hela_union = pd.read_csv(ROOT / "data/processed/HeLa_union.csv", dtype=str, keep_default_na=False)
    verify_union_means("HEK", hek_union, ["m6a_file1_s23", "m6a_file2"])
    verify_union_means("HeLa", hela_union, ["m6a_file1_sheet1", "m6a_file3"])
    compare("SAC HEK sheets 2+3", hek_expected, source_column(hek_union, "m6a_file1_s23"))
    compare("SAC HeLa sheet 1", hela_expected, source_column(hela_union, "m6a_file1_sheet1"))
    print(f"SAC rows: HEK polyA={poly_rows}, HEK ribo-={ribo_rows}, "
          f"combined unique={len(hek_expected)}, cross-sheet overlap={poly_rows + ribo_rows - len(hek_expected)}, "
          f"HeLa polyA={hela_rows}")
    glori = openpyxl.load_workbook(
        ROOT / "GLORI_HEK293T_mRNA.xlsx", read_only=True, data_only=True
    )
    glori_expected = {}
    rows = glori.active.iter_rows(values_only=True)
    next(rows)
    for chrom, pos, strand, _gene, _cov1, _cov2, rep1, rep2, *_ in rows:
        site = f"{chrom}_{int(pos)}_{strand}"
        if site in glori_expected:
            raise AssertionError(f"Duplicate GLORI site: {site}")
        a, b = dec(rep1), dec(rep2)
        if not (Decimal(0) <= a <= 1 and Decimal(0) <= b <= 1):
            raise ValueError(f"GLORI fraction out of range: {site}")
        glori_expected[site] = (a + b) * 50
    compare("GLORI HEK two-replicate mean ×100", glori_expected,
            source_column(hek_union, "m6a_file2"))
    print("PASS: all selected SAC and GLORI source positions and values match the union columns")


if __name__ == "__main__":
    main()
